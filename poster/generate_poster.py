"""
Maritime Grids — Final A1 Landscape Academic Poster Generator v2
St. Vincent Pallotti College of Engineering & Technology, Nagpur
Department of Computer Science and Engineering (Data Science)
Session 2025–26

Output:
  C:\maritime grids\poster\Maritime_Grids_Final_Poster_A1_Landscape.pptx

Reference style: light-blue institutional header, three-column academic layout,
navy/white/pale-blue palette, Times New Roman title, Calibri body.
"""

import os
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ─────────────────────────────────────────────────────────────────────────────
# SIZES  (A1 Landscape = 841 × 594 mm)
# ─────────────────────────────────────────────────────────────────────────────
A1_W = 84.1
A1_H = 59.4

# ─────────────────────────────────────────────────────────────────────────────
# COLOUR PALETTE
# ─────────────────────────────────────────────────────────────────────────────
NAVY      = RGBColor(0x0D, 0x2C, 0x5C)
BLUE_HDR  = RGBColor(0xC5, 0xDB, 0xEF)   # pale-blue header
ACCENT    = RGBColor(0x17, 0x55, 0x9A)   # section-title stripe
OFFWHITE  = RGBColor(0xF5, 0xF7, 0xFA)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
LTGRAY    = RGBColor(0xD0, 0xD8, 0xE4)
DARKGRAY  = RGBColor(0x2A, 0x2A, 0x3A)
GREEN_D   = RGBColor(0x12, 0x6B, 0x3A)
GREEN_L   = RGBColor(0xE6, 0xF5,0xED)
ORANGE    = RGBColor(0xC8, 0x60, 0x00)
ORANGE_L  = RGBColor(0xFF, 0xF4, 0xE0)
RED       = RGBColor(0xBB, 0x00, 0x00)
RED_L     = RGBColor(0xFF, 0xEB, 0xEB)
BLUE_L    = RGBColor(0xE8, 0xF3, 0xFD)
BLUE_M    = RGBColor(0xD0, 0xE8, 0xFF)

# Paths
ROOT    = r"C:\maritime grids"
POSTER  = os.path.join(ROOT, "poster")
FIGS    = os.path.join(ROOT, "docs", "figures")
OUT     = os.path.join(POSTER, "Maritime_Grids_Final_Poster_A1_Landscape.pptx")


# ─────────────────────────────────────────────────────────────────────────────
# PRIMITIVE HELPERS
# ─────────────────────────────────────────────────────────────────────────────
def c(v): return Cm(v)


def rect(slide, x, y, w, h, fill=None, border=None, bpt=0.6):
    from pptx.util import Pt as _Pt
    s = slide.shapes.add_shape(1, c(x), c(y), c(w), c(h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    else:
        s.fill.background()
    if border:
        s.line.color.rgb = border; s.line.width = _Pt(bpt)
    else:
        s.line.fill.background()
    return s


def txb(slide, x, y, w, h, text, sz=7, bold=False, italic=False,
        col=DARKGRAY, align=PP_ALIGN.LEFT, wrap=True, fn="Calibri"):
    from pptx.util import Pt as _Pt
    tb = slide.shapes.add_textbox(c(x), c(y), c(w), c(h))
    tb.word_wrap = wrap
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = _Pt(sz)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = col
    r.font.name = fn
    return tb


def mtxb(slide, x, y, w, h, lines, szs, bolds=None, cols=None,
         align=PP_ALIGN.LEFT, fn="Calibri", spc=0):
    from pptx.util import Pt as _Pt
    tb = slide.shapes.add_textbox(c(x), c(y), c(w), c(h))
    tb.word_wrap = True
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = _Pt(spc)
        r = p.add_run()
        r.text = line
        r.font.size = _Pt(szs[i] if isinstance(szs, list) else szs)
        r.font.bold  = bolds[i] if bolds else False
        r.font.color.rgb = cols[i] if cols else DARKGRAY
        r.font.name  = fn
    return tb


def img(slide, path, x, y, w, h):
    if os.path.exists(path):
        slide.shapes.add_picture(path, c(x), c(y), c(w), c(h))
    else:
        rect(slide, x, y, w, h, fill=LTGRAY, border=ACCENT)
        txb(slide, x+0.1, y+h/2-0.2, w-0.2, 0.4,
            f"[{os.path.basename(path)}]", sz=5.5, col=DARKGRAY, align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────────────────────────────────────
# SECTION BOX: coloured title bar + body text
# ─────────────────────────────────────────────────────────────────────────────
def section(slide, x, y, w, h, title, body_lines,
            tbg=ACCENT, tfg=WHITE, bbg=OFFWHITE, bdr=LTGRAY,
            tsz=8.5, bsz=6.4, th=0.88, indent=False, fn="Calibri"):
    rect(slide, x, y, w, h, fill=bbg, border=bdr, bpt=0.5)
    rect(slide, x, y, w, th, fill=tbg, border=None)
    txb(slide, x+0.12, y+0.06, w-0.24, th-0.1, title,
        sz=tsz, bold=True, col=tfg, align=PP_ALIGN.CENTER, fn=fn)
    if body_lines:
        body = "\n".join(body_lines)
        txb(slide, x+0.15, y+th+0.07, w-0.3, h-th-0.12,
            body, sz=bsz, col=DARKGRAY, wrap=True, fn=fn)


# ─────────────────────────────────────────────────────────────────────────────
# ARCH BOX (rounded look inside architecture diagram)
# ─────────────────────────────────────────────────────────────────────────────
def abox(slide, x, y, w, h, title, items, fill=WHITE, tcol=NAVY,
         bdr=ACCENT, tsz=7.2, isz=6.0, fn="Calibri"):
    rect(slide, x, y, w, h, fill=fill, border=bdr, bpt=0.75)
    txb(slide, x+0.1, y+0.05, w-0.2, 0.48,
        title, sz=tsz, bold=True, col=tcol, align=PP_ALIGN.CENTER, fn=fn)
    body = "\n".join(["• "+i for i in items])
    txb(slide, x+0.15, y+0.52, w-0.3, h-0.6,
        body, sz=isz, col=DARKGRAY, wrap=True, fn=fn)


def arrow(slide, cx, top, gap=0.22):
    txb(slide, cx-0.5, top+0.01, 1.0, gap-0.01,
        "▼", sz=8, bold=True, col=ACCENT, align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────────────────────────────────────
# MAIN BUILD
# ─────────────────────────────────────────────────────────────────────────────
def build():
    prs = Presentation()
    prs.slide_width  = c(A1_W)
    prs.slide_height = c(A1_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # ── Full white background ────────────────────────────────────────────────
    rect(slide, 0, 0, A1_W, A1_H, fill=WHITE)

    # ════════════════════════════════════════════════════════════════════════
    # HEADER  (height 7.2 cm)
    # ════════════════════════════════════════════════════════════════════════
    HH = 7.2
    rect(slide, 0, 0, A1_W, HH, fill=BLUE_HDR, border=NAVY, bpt=1.2)

    # College logo box
    rect(slide, 0.4, 0.35, 5.6, 5.8, fill=WHITE, border=NAVY, bpt=0.5)
    txb(slide, 0.5, 0.55, 5.4, 5.4,
        "St. Vincent\nPallotti\nCollege\nof Engineering\n& Technology\n\nNagpur",
        sz=6.8, bold=True, col=NAVY, align=PP_ALIGN.CENTER, fn="Times New Roman")

    # ARISE & SHINE badge
    rect(slide, 78.0, 0.35, 5.6, 5.8, fill=WHITE, border=NAVY, bpt=0.5)
    txb(slide, 78.1, 0.55, 5.4, 5.4,
        "ARISE\n&\nSHINE",
        sz=9, bold=True, col=NAVY, align=PP_ALIGN.CENTER, fn="Times New Roman")

    # Centre titles
    mtxb(slide, 7.5, 0.15, 69.0, 6.8,
         [
             "St. Vincent Pallotti College of Engineering & Technology, Nagpur",
             "Department of Computer Science and Engineering (Data Science)",
             "Academic Session 2025–26",
             " ",
             "MARITIME GRIDS",
             "Integrated Maritime Voyage Simulation, Energy Modeling, Optimization and Decision Support",
         ],
         szs=[11, 10, 9.5, 3, 24, 10.5],
         bolds=[True, True, False, False, True, False],
         cols=[NAVY, NAVY, NAVY, WHITE, NAVY, ACCENT],
         align=PP_ALIGN.CENTER, fn="Times New Roman")

    # Guide + team
    txb(slide, 7.5, 5.75, 69.0, 1.3,
        "Guide: Dr. Manish Gudadhe                    "
        "Team:  Chinmayakumar Shegokar  |  Prajakta Timande  |  Tanushri Tonde  |  Vuha Modala",
        sz=9, col=NAVY, align=PP_ALIGN.CENTER, fn="Calibri")

    # Thin separator line
    rect(slide, 0, HH-0.06, A1_W, 0.06, fill=NAVY)

    # ════════════════════════════════════════════════════════════════════════
    # THREE-COLUMN BODY
    # ════════════════════════════════════════════════════════════════════════
    TOP  = HH + 0.22
    BH   = A1_H - TOP - 0.2          # total body height
    CW   = (A1_W - 0.7) / 3.0        # column width
    C1   = 0.22
    C2   = C1 + CW + 0.13
    C3   = C2 + CW + 0.13

    # ════════════════════════════════════════════════════════════════════════
    # LEFT COLUMN
    # ════════════════════════════════════════════════════════════════════════
    y1 = TOP

    # ── Abstract ────────────────────────────────────────────────────────────
    ah = 7.2
    section(slide, C1, y1, CW, ah, "ABSTRACT", [],
            tsz=8.8, th=0.88)
    txb(slide, C1+0.15, y1+1.0, CW-0.3, ah-1.12,
        "Maritime voyage operations require simultaneous management of vessel "
        "dynamics, route geometry, environmental conditions (wind, waves, current), "
        "fuel and energy consumption, battery behavior, and regulatory compliance. "
        "These are typically handled in isolation, creating fragmented decision-making.\n\n"
        "Maritime Grids provides an integrated engineering platform combining "
        "physics-based voyage simulation, environmental forecast processing, fuel and "
        "hybrid energy modeling, multi-objective optimization, IMO/CII emissions "
        "analysis, and human-in-the-loop decision support via the Signal Agent.\n\n"
        "A Native Virtual Maritime Simulator demonstrates the pipeline end-to-end on the "
        "Escravos Terminal → Lagos Port route (250 NM), verified through 975 automated tests.",
        sz=6.7, col=DARKGRAY, fn="Calibri")
    y1 += ah + 0.2

    # ── Problem Statement ───────────────────────────────────────────────────
    ph = 6.2
    section(slide, C1, y1, CW, ph, "PROBLEM STATEMENT",
            ["Maritime voyage planning requires simultaneous",
             "consideration of:",
             "  • Vessel characteristics & route geometry",
             "  • Wind, waves & ocean current forces",
             "  • Fuel consumption & hybrid energy behavior",
             "  • Battery SOC / SOH evolution",
             "  • Voyage duration, ETA & progress",
             "  • CO₂ emissions & CII/regulatory obligations",
             "  • Economic & engineering trade-offs",
             "",
             "These aspects are often handled separately.",
             "Maritime Grids integrates them into one",
             "unified engineering decision platform."],
            tsz=8.8, bsz=6.4)
    y1 += ph + 0.2

    # ── Objectives ──────────────────────────────────────────────────────────
    oh = 8.6
    section(slide, C1, y1, CW, oh, "OBJECTIVES",
            ["O1  Integrate maritime, environmental, vessel & route",
             "     data into a common engineering platform.",
             "O2  Model vessel movement using physics-based",
             "     maritime simulation.",
             "O3  Estimate fuel, propulsion power & hybrid energy",
             "     (battery + solar) behavior.",
             "O4  Evaluate vessel sizing, speed & energy trade-offs.",
             "",
             "O5  Perform optimization, sensitivity, break-even &",
             "     fleet-level engineering analysis.",
             "O6  Analyze emissions & CII/regulatory implications",
             "     of voyage scenarios.",
             "O7  Provide human-in-the-loop operational decision",
             "     support via Signal Agent & Driver Dashboard."],
            tsz=8.8, bsz=6.4)
    y1 += oh + 0.2

    # ── Environmental Intelligence ──────────────────────────────────────────
    eh = 8.4
    section(slide, C1, y1, CW, eh, "ENVIRONMENTAL INTELLIGENCE",
            ["Environmental inputs processed by Maritime Grids:",
             "  • Wind Speed & Direction  (m/s, degrees)",
             "  • Significant Wave Height & Period",
             "  • Wave Direction",
             "  • Ocean Current Speed & Direction",
             "  • Air Temperature & Relative Humidity",
             "  • Atmospheric Pressure",
             "  • Solar Irradiance / GHI  (Climatological Est.)",
             "",
             "Provider: Open-Meteo Forecast API + Mock Provider",
             "Route Lookahead: Future waypoints evaluated",
             "ahead of vessel position for adverse conditions.",
             "",
             "Threshold logic:",
             "  Wind > 25 kn → MET_WIND_HIGH_001",
             "  Waves > 2.5 m → MET_SEA_ROUGH_002",
             "  Ahead adverse → MET_AHEAD_ADVERSE_001"],
            tsz=8.8, bsz=6.2)

    # Provenance mini boxes inside env section
    px = C1 + 0.15
    py = y1 + eh - 2.15
    phw = (CW - 0.3) / 2 - 0.06
    rect(slide, px, py, phw, 2.0, fill=BLUE_L, border=ACCENT, bpt=0.5)
    mtxb(slide, px+0.08, py+0.12, phw-0.16, 1.75,
         ["ENVIRONMENTAL INPUT", "FORECAST / MODEL DATA",
          "Weather · Waves · Current"],
         szs=[6.5, 6.0, 5.8], bolds=[True, False, False],
         cols=[ACCENT, DARKGRAY, DARKGRAY], align=PP_ALIGN.CENTER)

    px2 = px + phw + 0.12
    rect(slide, px2, py, phw, 2.0, fill=GREEN_L, border=GREEN_D, bpt=0.5)
    mtxb(slide, px2+0.08, py+0.12, phw-0.16, 1.75,
         ["SIMULATOR OUTPUT", "SIMULATED / DERIVED",
          "Position · SOG · Fuel · SOC"],
         szs=[6.5, 6.0, 5.8], bolds=[True, False, False],
         cols=[GREEN_D, DARKGRAY, DARKGRAY], align=PP_ALIGN.CENTER)
    y1 += eh + 0.2

    # ── Maritime Physics ────────────────────────────────────────────────────
    mh = 6.8
    section(slide, C1, y1, CW, mh, "MARITIME PHYSICS MODEL",
            ["Wind + Waves + Current",
             "        ↓   Environmental Forces Applied",
             "Hydrodynamic Resistance Calculated",
             "        ↓",
             "Required Propulsion Power Derived",
             "        ↓",
             "Ship Speed / SOG Determined",
             "        ↓",
             "Fuel Rate & Energy Consumption",
             "",
             "The simulator applies physics to calculate vessel",
             "behavior — not a simple icon-moving simulator.",
             "Each step: Resistance → Power → Fuel → SOC."],
            tsz=8.8, bsz=6.4)
    y1 += mh + 0.2

    # ── Fuel · Battery · Solar ──────────────────────────────────────────────
    fbs_h = BH - (y1 - TOP)
    section(slide, C1, y1, CW, max(fbs_h, 9.0), "FUEL · BATTERY · SOLAR",
            ["FUEL  (Physics / Model-Derived):",
             "  Shaft/propulsion power → fuel rate",
             "  Fuel consumed & remaining tracked",
             "  Speed-dependent & resistance-adjusted",
             "  XGBoost fuel model (R² = 0.9513)",
             "",
             "BATTERY  (Simulated / Model-Derived):",
             "  SOC — 2-layer LSTM  (R² = 0.9089 on B0007)",
             "  SOH — HistGradientBoosting  (Test R² = 0.6308)",
             "  Usable energy = Capacity × SOH",
             "  Pack topology: 96S×nP per vessel profile",
             "",
             "SOLAR  (Climatological / Scenario Estimate):",
             "  NASA POWER monthly GHI data",
             "  Solar area × panel efficiency × GHI",
             "  Auxiliary hotel-load coverage modeled",
             "  Labeled as scenario estimate, not live data"],
            tsz=8.8, bsz=6.3)

    # ════════════════════════════════════════════════════════════════════════
    # CENTRE COLUMN
    # ════════════════════════════════════════════════════════════════════════
    y2 = TOP

    # ── Core Architecture Diagram ────────────────────────────────────────────
    arch_h = 35.5
    rect(slide, C2, y2, CW, arch_h, fill=OFFWHITE, border=LTGRAY, bpt=0.4)
    rect(slide, C2, y2, CW, 0.88, fill=NAVY)
    txb(slide, C2, y2+0.06, CW, 0.78,
        "COMPLETE SYSTEM ARCHITECTURE",
        sz=9.5, bold=True, col=WHITE, align=PP_ALIGN.CENTER, fn="Calibri")

    bx  = C2 + 0.45
    bw  = CW - 0.9
    bh  = 1.72
    gap = 0.21
    ay  = y2 + 1.0

    def draw_arch_box(ay, title, items, fill=WHITE, tcol=NAVY):
        abox(slide, bx, ay, bw, bh, title, items, fill=fill, tcol=tcol,
             tsz=7.0, isz=5.9)
        return ay + bh

    def draw_arrow(ay):
        arrow(slide, C2 + CW/2, ay, gap)
        return ay + gap

    ay = draw_arch_box(ay, "MARITIME & ENVIRONMENTAL DATA",
        ["Weather / Wind / Waves / Ocean Current",
         "Temp · Humidity · Pressure · Solar/GHI",
         "Route Data · Vessel Parameters"],
        fill=BLUE_L)
    ay = draw_arrow(ay)

    ay = draw_arch_box(ay, "DATA ENGINEERING",
        ["Ingestion · QC · Validation",
         "Spatial Processing",
         "Temporal / Route Alignment"],
        fill=RGBColor(0xEE,0xF5,0xFF))
    ay = draw_arrow(ay)

    ay = draw_arch_box(ay, "ENVIRONMENTAL INTELLIGENCE",
        ["Weather Forecast · Route Lookahead",
         "Adverse Weather Detection",
         "Open-Meteo API + Mock Provider"],
        fill=BLUE_L)
    ay = draw_arrow(ay)

    ay = draw_arch_box(ay, "MARITIME PHYSICS",
        ["Vessel Dynamics · Resistance",
         "Environmental Forces",
         "Propulsion Power · Speed"],
        fill=RGBColor(0xEE,0xF5,0xFF))
    ay = draw_arrow(ay)

    ay = draw_arch_box(ay, "ENERGY & FUEL MODEL",
        ["Fuel Consumption — XGBoost (Model B)",
         "Battery SOC — LSTM (Model C-SOC)",
         "Battery SOH — HGB · Solar / GHI"],
        fill=BLUE_L)
    ay = draw_arrow(ay)

    ay = draw_arch_box(ay, "VOYAGE SIMULATOR",
        ["Position · SOG / STW · Heading",
         "Distance · Progress · ETA",
         "Native Virtual Maritime Simulator"],
        fill=BLUE_M, tcol=ACCENT)
    ay = draw_arrow(ay)

    # Split → two boxes
    split_y = ay
    hbw = (bw - 0.15) / 2

    # Engineering Workbench (left half)
    rect(slide, bx, split_y, hbw, 3.5, fill=GREEN_L, border=GREEN_D, bpt=0.75)
    txb(slide, bx+0.05, split_y+0.08, hbw-0.1, 0.52,
        "ENGINEERING\nWORKBENCH", sz=6.5, bold=True, col=GREEN_D,
        align=PP_ALIGN.CENTER)
    txb(slide, bx+0.1, split_y+0.65, hbw-0.2, 2.8,
        "• Voyage Simulation\n"
        "• Sizing Evaluation\n"
        "• Sizing Optimization\n"
        "• Speed Co-Optimization\n"
        "• Sensitivity Analysis\n"
        "• Break-Even Analysis\n"
        "• 2D Grid Analysis\n"
        "• Fleet Portfolio\n"
        "• CII Trajectory\n"
        "• Scenario Dossier",
        sz=5.7, col=DARKGRAY)

    # Signal Agent (right half)
    sx = bx + hbw + 0.15
    rect(slide, sx, split_y, hbw, 3.5, fill=ORANGE_L, border=ORANGE, bpt=0.75)
    txb(slide, sx+0.05, split_y+0.08, hbw-0.1, 0.52,
        "SIGNAL AGENT\nDECISION SUPPORT", sz=6.5, bold=True, col=ORANGE,
        align=PP_ALIGN.CENTER)
    txb(slide, sx+0.1, split_y+0.65, hbw-0.2, 2.8,
        "• Weather Alerts\n"
        "• Fuel Alerts\n"
        "• Battery Alerts\n"
        "• CII / Regulatory\n"
        "• Decision Not Ready\n"
        "• Human Review\n"
        "• human_review\n"
        "  required = True",
        sz=5.7, col=DARKGRAY)

    ay = split_y + 3.65

    txb(slide, C2 + CW/2 - 0.5, ay, 1.0, gap,
        "▼", sz=8, bold=True, col=NAVY, align=PP_ALIGN.CENTER)
    ay += gap + 0.05

    # Dashboard box
    rect(slide, bx, ay, bw, 1.45, fill=BLUE_M, border=NAVY, bpt=1.0)
    txb(slide, bx+0.1, ay+0.1, bw-0.2, 0.65,
        "DRIVER DASHBOARD  +  ENGINEERING WORKBENCH",
        sz=7.5, bold=True, col=NAVY, align=PP_ALIGN.CENTER)
    txb(slide, bx+0.1, ay+0.75, bw-0.2, 0.6,
        "▼  HUMAN DECISION  ▼",
        sz=7, bold=True, col=ACCENT, align=PP_ALIGN.CENTER)

    y2 += arch_h + 0.25

    # ── Model metrics ────────────────────────────────────────────────────────
    mm_h = 5.6
    section(slide, C2, y2, CW, mm_h, "ML / ENGINEERING MODELS — VERIFIED METRICS",
            ["Model B — Fuel Consumption  (XGBoost, ship-holdout):",
             "   R² = 0.9513   RMSE = 1092 L   MAE = 662 L",
             "   CO₂ RMSE = 3310 kg, R² = 0.9423",
             "",
             "Model C-SOC — Battery State of Charge  (2-layer LSTM):",
             "   B0007 Unseen-Cell Test:  R² = 0.9089  RMSE = 0.0678",
             "",
             "Model C-SOH — Battery State of Health  (HistGradBoost):",
             "   Val R² = 0.9028   Test RMSE = 0.0397  (B0007)",
             "",
             "Model A — Weather Classifier:",
             "   Honest result: No ML model exceeded majority-class",
             "   baseline (34.72%). Majority class selected."],
            tsz=8.5, bsz=6.3)
    y2 += mm_h + 0.22

    # ── Two model charts side by side ───────────────────────────────────────
    cw2 = (CW - 0.5) / 2
    ch  = 7.8
    img(slide, os.path.join(FIGS, "battery_soc_actual_vs_predicted.png"),
        C2+0.12, y2, cw2, ch)
    img(slide, os.path.join(FIGS, "fuel_actual_vs_predicted.png"),
        C2+0.22+cw2, y2, cw2, ch)
    y2 += ch + 0.18

    # ── SOH degradation chart ────────────────────────────────────────────────
    soh_h = BH - (y2 - TOP)
    img(slide, os.path.join(FIGS, "battery_soh_degradation.png"),
        C2+0.22, y2, CW-0.44, max(soh_h, 3.5))

    # ════════════════════════════════════════════════════════════════════════
    # RIGHT COLUMN
    # ════════════════════════════════════════════════════════════════════════
    y3 = TOP

    # ── Project Outcomes ────────────────────────────────────────────────────
    po_h = 8.0
    section(slide, C3, y3, CW, po_h, "PROJECT OUTCOMES",
            ["PO1 — Integrated Maritime Data",
             "  Environmental, route & vessel data unified.",
             "",
             "PO2 — Physics-Based Voyage Simulation",
             "  Virtual vessel coupled with environmental",
             "  conditions, dynamics & energy models.",
             "",
             "PO3 — Engineering Intelligence",
             "  Sizing, speed, sensitivity, break-even,",
             "  grid & fleet-level analyses.",
             "",
             "PO4 — Regulatory & Emissions Analysis",
             "  Voyage → CII trajectory via MEPC.353/354(78).",
             "",
             "PO5 — Operational Decision Support",
             "  Signal Agent: weather, fuel, battery, CII",
             "  advisories — human review required.",
             "",
             "PO6 — Verified Integrated Platform",
             "  975 automated tests + browser acceptance."],
            tsz=8.8, bsz=6.3)
    y3 += po_h + 0.2

    # ── Verification / Results ──────────────────────────────────────────────
    rv_h = 7.0
    rect(slide, C3, y3, CW, rv_h, fill=GREEN_L, border=GREEN_D, bpt=0.75)
    rect(slide, C3, y3, CW, 0.88, fill=GREEN_D)
    txb(slide, C3, y3+0.06, CW, 0.78, "VERIFICATION & RESULTS",
        sz=8.8, bold=True, col=WHITE, align=PP_ALIGN.CENTER)

    txb(slide, C3+0.3, y3+0.98, CW-0.6, 1.15,
        "975 TESTS PASSED",
        sz=18, bold=True, col=GREEN_D, align=PP_ALIGN.CENTER, fn="Calibri")

    hw = (CW - 0.5) / 2
    rect(slide, C3+0.2, y3+2.2, hw, 0.7, fill=WHITE, border=GREEN_D, bpt=0.5)
    txb(slide, C3+0.2, y3+2.25, hw, 0.65,
        "Frontend  65 / 65", sz=8, bold=True, col=ACCENT, align=PP_ALIGN.CENTER)

    rect(slide, C3+0.3+hw, y3+2.2, hw, 0.7, fill=WHITE, border=GREEN_D, bpt=0.5)
    txb(slide, C3+0.3+hw, y3+2.25, hw, 0.65,
        "Backend  910 / 910", sz=8, bold=True, col=ACCENT, align=PP_ALIGN.CENTER)

    txb(slide, C3+0.18, y3+3.05, CW-0.36, rv_h-3.2,
        "✓ Native simulator movement verified\n"
        "✓ Route progress & distance accumulation\n"
        "✓ ETA calculation verified\n"
        "✓ Weather synchronization verified\n"
        "✓ Signal Agent advisory chain verified\n"
        "✓ Engine provenance integrity verified\n"
        "✓ Map & telemetry synchronization\n"
        "✓ Fuel / battery evolution verified\n"
        "✓ Pause / reset cycle verified\n"
        "✓ Engineering Workbench preserved\n"
        "✓ 6,228 raw data files validated\n"
        "✓ 0 actuator commands confirmed",
        sz=6.3, col=DARKGRAY, fn="Calibri")
    y3 += rv_h + 0.2

    # ── Engineering Workbench panel ─────────────────────────────────────────
    ew_h = 8.8
    rect(slide, C3, y3, CW, ew_h, fill=GREEN_L, border=GREEN_D, bpt=0.75)
    rect(slide, C3, y3, CW, 0.88, fill=GREEN_D)
    txb(slide, C3, y3+0.06, CW, 0.78, "ENGINEERING WORKBENCH",
        sz=8.8, bold=True, col=WHITE, align=PP_ALIGN.CENTER)

    mods = ["Voyage Simulation", "Sizing Evaluation",
            "Sizing Optimization", "Speed Co-Optimization",
            "Sensitivity Analysis", "Break-Even Analysis",
            "2D Grid Analysis", "Fleet Portfolio",
            "CII Trajectory", "Scenario Dossier / Export"]
    mx = C3 + 0.22
    mw = CW - 0.44
    mhh = 0.60
    my = y3 + 0.98
    for i, mod in enumerate(mods):
        bg = RGBColor(0xE0,0xF4,0xE0) if i % 2 == 0 else RGBColor(0xCC,0xEC,0xCC)
        rect(slide, mx, my, mw, mhh, fill=bg, border=GREEN_D, bpt=0.3)
        txb(slide, mx+0.12, my+0.06, mw-0.24, mhh-0.1,
            mod, sz=6.7, bold=(i == 0), col=DARKGRAY)
        my += mhh + 0.03

    txb(slide, C3+0.15, y3+ew_h-0.72, CW-0.3, 0.65,
        "Simulation → Analysis → Optimization → Decision",
        sz=6.7, bold=True, col=GREEN_D, align=PP_ALIGN.CENTER)
    y3 += ew_h + 0.2

    # ── Signal Agent ────────────────────────────────────────────────────────
    sa_h = 7.0
    rect(slide, C3, y3, CW, sa_h, fill=ORANGE_L, border=ORANGE, bpt=0.75)
    rect(slide, C3, y3, CW, 0.88, fill=ORANGE)
    txb(slide, C3, y3+0.06, CW, 0.78, "SIGNAL AGENT — HUMAN-IN-THE-LOOP",
        sz=8.8, bold=True, col=WHITE, align=PP_ALIGN.CENTER)

    sa_items = [
        ("Vessel State + Weather + Fuel + Battery", BLUE_L),
        ("+ Route Lookahead + Regulatory State", BLUE_L),
        ("▼  SIGNAL AGENT  ▼", None),
        ("NORMAL", GREEN_L),
        ("WEATHER ALERT", ORANGE_L),
        ("FUEL ALERT", ORANGE_L),
        ("BATTERY ALERT", ORANGE_L),
        ("REGULATORY / CII ALERT", ORANGE_L),
        ("▼  HUMAN REVIEW  ▼", None),
    ]
    sfy = y3 + 0.95
    sfh = 0.52
    for txt, fill_ in sa_items:
        if fill_ is None:
            txb(slide, C3+0.4, sfy, CW-0.8, sfh,
                txt, sz=7, bold=True, col=ORANGE, align=PP_ALIGN.CENTER)
        else:
            rect(slide, C3+0.28, sfy, CW-0.56, sfh-0.04, fill=fill_, border=ORANGE, bpt=0.3)
            txb(slide, C3+0.35, sfy+0.03, CW-0.7, sfh-0.1,
                txt, sz=6.3, col=DARKGRAY, align=PP_ALIGN.CENTER)
        sfy += sfh + 0.04

    txb(slide, C3+0.15, sfy+0.06, CW-0.3, 0.65,
        "Signal Agent does NOT control the vessel.\nAll advisories require human review.",
        sz=6.2, bold=True, col=ORANGE, align=PP_ALIGN.CENTER)
    y3 += sa_h + 0.2

    # ── IMO / CII / Emissions ───────────────────────────────────────────────
    cii_h = 6.2
    section(slide, C3, y3, CW, cii_h, "IMO / CII / EMISSIONS ANALYSIS",
            ["Voyage + Speed + Fuel Consumption",
             "       ↓  CO₂ Emissions Derived",
             "CII = CO₂ × 1000 / (DWT × distance NM)",
             "       ↓  Rating: A / B / C / D / E",
             "MARPOL Annex VI Regulation 28:",
             "  Statutory Z% Schedule 2023–2026:",
             "  5% → 7% → 9% → 11%",
             "Reference Line: MEPC.353(78)",
             "Rating Bands: MEPC.354(78)",
             "Multi-year CII trajectory simulation.",
             "Clean-energy assisted compliance extension.",
             "No official regulatory certification claimed."],
            tsz=8.8, bsz=6.3)
    y3 += cii_h + 0.2

    # ── Telemetry + Bridge Command ──────────────────────────────────────────
    tel_h = 4.8
    section(slide, C3, y3, CW, tel_h, "TELEMETRY & BRIDGE COMMAND",
            ["Bridge Command 5.11.0  (Read-Only Standby)",
             "   NMEA 0183 UDP → Port 10110",
             "   $GPRMC · $GPGGA · $GPHDT · $GPVTG",
             "   NMEA Adapter → CanonicalTelemetryPayload",
             "   Telemetry Gateway → SSE/REST → Dashboard",
             "",
             "Live evidence: SOG=5.0 kn · COG=90° confirmed",
             "Native Simulator = Primary Demonstration",
             "Bridge Command = Optional / Read-Only Standby",
             "Port 10111: 0 connections (read-only verified)"],
            tsz=8.8, bsz=6.3)
    y3 += tel_h + 0.2

    # ── Safety Boundary ─────────────────────────────────────────────────────
    sf_h = 4.8
    rect(slide, C3, y3, CW, sf_h, fill=RED_L, border=RED, bpt=1.0)
    rect(slide, C3, y3, CW, 0.88, fill=RED)
    txb(slide, C3, y3+0.06, CW, 0.78, "⚠  SAFETY BOUNDARY",
        sz=8.8, bold=True, col=WHITE, align=PP_ALIGN.CENTER)
    mtxb(slide, C3+0.22, y3+0.98, CW-0.44, sf_h-1.1,
         ["✗  0 External Actuators",
          "✗  0 External Control Commands",
          "✓  Read-Only Operational Integration",
          "✓  Human Review Required for All Advisories",
          "✗  No Autonomous Steering",
          "✗  No Autonomous Throttle",
          "✗  No Autonomous Engine Control"],
         szs=[7.3]*7, bolds=[True]*7,
         cols=[RED, RED, GREEN_D, GREEN_D, RED, RED, RED])
    y3 += sf_h + 0.2

    # ── Applications ────────────────────────────────────────────────────────
    ap_h = 4.9
    section(slide, C3, y3, CW, ap_h, "APPLICATIONS",
            ["• Voyage planning & vessel performance analysis",
             "• Fuel optimization",
             "• Hybrid energy planning (battery + solar)",
             "• Weather-aware voyage analysis",
             "• Vessel sizing engineering",
             "• Speed & fleet optimization",
             "• Emissions / CII scenario analysis",
             "• Operational decision support"],
            tsz=8.8, bsz=6.4)
    y3 += ap_h + 0.2

    # ── Future Scope ─────────────────────────────────────────────────────────
    fs_h = BH - (y3 - TOP)
    section(slide, C3, y3, CW, max(fs_h, 3.5), "FUTURE SCOPE",
            ["• Larger route library",
             "• Additional vessel profiles",
             "• More real-world telemetry sources",
             "• Advanced ML weather forecasting",
             "• Expanded environmental datasets",
             "• Additional optimization objectives",
             "• Larger fleet-scale simulations",
             "• Additional regulatory scenarios",
             "• Real-time AIS integration (read-only)"],
            tsz=8.8, bsz=6.4)

    # ════════════════════════════════════════════════════════════════════════
    # NATIVE SIMULATOR PANEL — spans left + centre columns bottom strip
    # ════════════════════════════════════════════════════════════════════════
    # Calculate remaining left-column space
    sim_top = y1      # y1 currently points to start of fuel/battery section
    # Instead place it under the fuel section
    sim_panel_top = TOP + BH - 12.5
    sim_panel_h   = 12.5 - 0.05
    sim_panel_w   = C2 + CW - C1 - 0.1

    # Override: put the simulator section in the lower portion of column 1
    # (we already filled col 1 fully; this is shown inside fuel section space which
    # was designed to fill remaining height). Skip this in favor of the existing layout
    # which has the fuel section fill remaining height.

    # ════════════════════════════════════════════════════════════════════════
    # ROUTE WEATHER LOOKAHEAD (bottom of centre, below charts)
    # ════════════════════════════════════════════════════════════════════════
    # (Already included as part of the arch diagram + env intelligence section)

    # ════════════════════════════════════════════════════════════════════════
    # BOTTOM FOOTER STRIP
    # ════════════════════════════════════════════════════════════════════════
    footer_y = A1_H - 0.5
    rect(slide, 0, footer_y, A1_W, 0.5, fill=NAVY)
    txb(slide, 0.3, footer_y+0.05, A1_W-0.6, 0.4,
        "Maritime Grids  |  Integrated Maritime Simulation & Engineering Intelligence Platform  |  "
        "St. Vincent Pallotti College of Engineering & Technology, Nagpur  |  2025–26",
        sz=6, col=WHITE, align=PP_ALIGN.CENTER, fn="Calibri")

    return prs


def main():
    print("Building Maritime Grids A1 Landscape Poster (v2)…")
    prs = build()
    prs.save(OUT)
    print(f"  ✓ PPTX saved → {OUT}")

    # PNG via PowerPoint COM if available
    try:
        import subprocess
        result = subprocess.run(
            ["python", "-c",
             f"""
import win32com.client, os
pptx = r'{OUT}'
png  = r'{os.path.join(POSTER, "Maritime_Grids_Final_Poster_A1_Landscape.png")}'
ppt = win32com.client.Dispatch('PowerPoint.Application')
ppt.Visible = True
pres = ppt.Presentations.Open(pptx, ReadOnly=True, WithWindow=False)
pres.Slides(1).Export(png, 'PNG', 4961, 3508)
pres.Close()
ppt.Quit()
print('PNG exported:', png)
"""],
            capture_output=True, text=True, timeout=90)
        if result.returncode == 0:
            print(f"  ✓ PNG exported")
        else:
            print("  ! PNG export via COM failed:", result.stderr[:200])
    except Exception as e:
        print(f"  ! PNG export skipped ({e})")

    print("Done.")


if __name__ == "__main__":
    main()
